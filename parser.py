import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

from selenium import webdriver
from selenium.webdriver.edge.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import (
    TimeoutException,
    StaleElementReferenceException,
)

from bs4 import BeautifulSoup
from tqdm import tqdm
import time
import json
from datetime import datetime
from collections import Counter

from pydantic import ValidationError
from openpyxl import Workbook

import db
from models import Product


# Конфигурация

CATEGORIES = {
    "Тело и гигиена": "https://www.yves-rocher.ru/collections/dlya-tela",
}

CARD_SELECTOR       = ".product-card"
NAME_SELECTOR       = 'a[class*="titleContainer"] div'
PRICE_SELECTOR      = 'span[class*="PriceLabel_price"]:not([class*="oldPrice"])'
OLD_PRICE_SELECTOR  = 'span[class*="oldPrice"]'
DISCOUNT_SELECTOR   = 'div[class*="priceDiscount"]'
IMAGE_SELECTOR      = "img.picture__img"
LINK_SELECTOR       = 'a[href]'


# Утилиты

def parse_price(price_str):
    if not price_str:
        return None
    cleaned = (price_str
               .replace("₽", "")
               .replace(" ", "")
               .replace("\xa0", "")
               .replace(",", ".")
               .strip())
    try:
        return float(cleaned)
    except ValueError:
        return None


def fix_url(link):
    if not link:
        return ""
    if link.startswith("/"):
        return "https://www.yves-rocher.ru" + link
    if not link.startswith("http"):
        return "https://www.yves-rocher.ru/" + link
    return link


def count_cards():
    return len(driver.find_elements(By.CSS_SELECTOR, CARD_SELECTOR))


def close_popups():
    """Удаляет рекламные попапы Flocktory, если они есть."""
    driver.execute_script("""
        const selectors = [
            'iframe.flocktory-widget',
            'iframe[title="Flocktory widget"]',
            'div[class*="flocktory"]'
        ];
        selectors.forEach(sel => {
            document.querySelectorAll(sel).forEach(el => el.remove());
        });
    """)


def click_load_more():
    """Кликает 'Показать ещё', пока кнопка доступна."""
    clicks = 0
    stale_retries = 0
    MAX_STALE_RETRIES = 5

    while True:
        try:
            btn = WebDriverWait(driver, 10).until(
                EC.element_to_be_clickable(
                    (By.XPATH, "//button[contains(., 'Показать ещё')]")
                )
            )
            driver.execute_script(
                "arguments[0].scrollIntoView({block: 'center', behavior: 'instant'});",
                btn
            )
            before = count_cards()
            driver.execute_script("arguments[0].click();", btn)
            clicks += 1

            try:
                WebDriverWait(driver, 40).until(
                    lambda d: count_cards() > before
                )
                print(f"  Клик {clicks}: загружено карточек — {count_cards()}")
                stale_retries = 0
                time.sleep(2)
            except TimeoutException:
                print(f"  Клик {clicks}: новых карточек не появилось, ждём ещё...")
                driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
                try:
                    WebDriverWait(driver, 20).until(
                        lambda d: count_cards() > before
                    )
                    print(f"  Клик {clicks}: загружено карточек — {count_cards()}")
                    stale_retries = 0
                    time.sleep(2)
                except TimeoutException:
                    print(f"  Клик {clicks}: карточки закончились")
                    break

        except TimeoutException:
            print(f"Кнопка 'Показать ещё' больше не найдена. Всего кликов: {clicks}")
            break

        except StaleElementReferenceException:
            stale_retries += 1
            print(f"  Предупреждение: элемент устарел, повтор {stale_retries}/{MAX_STALE_RETRIES}")
            if stale_retries >= MAX_STALE_RETRIES:
                print("  Слишком много устаревших элементов, останавливаюсь")
                break
            time.sleep(1)
            continue

        except Exception as e:
            print(f"  Предупреждение: {type(e).__name__}, пробую снова")
            time.sleep(1)
            continue

    return count_cards()


# Selenium

options = Options()
options.add_argument("--start-maximized")
options.add_argument("--disable-blink-features=AutomationControlled")

driver = webdriver.Edge(options=options)
wait = WebDriverWait(driver, 15)


# Подготовка БД

print("Подготовка базы данных...")
db.create_table_if_not_exists()
db.clear_products()
print("Таблица готова")


# Сбор данных

raw_products = []

for category_name, url in CATEGORIES.items():
    print(f"\nКатегория: {category_name}")
    print(f"URL: {url}")

    driver.get(url)
    print(f"Открыто: {driver.title}")

    try:
        wait.until(EC.presence_of_all_elements_located(
            (By.CSS_SELECTOR, CARD_SELECTOR)
        ))
    except TimeoutException:
        print("Карточки не появились, пропускаю категорию")
        continue

    total = click_load_more()
    print(f"Всего карточек в DOM: {total}")

    if total == 0:
        print("Карточек 0 — жду восстановления DOM...")
        try:
            WebDriverWait(driver, 15).until(lambda d: count_cards() > 0)
            print(f"DOM восстановлен, карточек: {count_cards()}")
        except TimeoutException:
            print("DOM не восстановился — пропускаю категорию")
            continue

    soup = BeautifulSoup(driver.page_source, "lxml")
    cards = soup.select(CARD_SELECTOR)
    print(f"BeautifulSoup нашёл: {len(cards)} карточек")

    for card in tqdm(cards, desc="Собираю карточки"):
        try:
            name_tag  = card.select_one(NAME_SELECTOR)
            price_tag = card.select_one(PRICE_SELECTOR)
            old_tag   = card.select_one(OLD_PRICE_SELECTOR)
            disc_tag  = card.select_one(DISCOUNT_SELECTOR)
            img_tag   = card.select_one(IMAGE_SELECTOR)
            link_tag  = card.select_one(LINK_SELECTOR)

            raw = {
                "category":  category_name,
                "name":      name_tag.get_text(strip=True)  if name_tag  else "",
                "price":     parse_price(price_tag.get_text(strip=True)) if price_tag else None,
                "old_price": parse_price(old_tag.get_text(strip=True))   if old_tag   else None,
                "discount":  disc_tag.get_text(strip=True) if disc_tag else "",
                "link":      fix_url(link_tag.get("href", "")) if link_tag else "",
                "image_url": img_tag.get("src", "") if img_tag else "",
                "brand":     {"name": "Yves Rocher", "country": "France", "is_organic": True},
            }
            raw_products.append(raw)
        except Exception as e:
            print(f"\n  Ошибка при сборе: {e}")

driver.quit()
print(f"\nСобрано сырых записей: {len(raw_products)}")


# Валидация

print("\nВалидация через Pydantic...")

good, warning, bad = [], [], []

for raw in tqdm(raw_products, desc="Валидирую"):
    try:
        product = Product(**raw)
        if product.is_warning():
            warning.append(product)
        else:
            good.append(product)
    except ValidationError as e:
        bad.append((raw, e.errors()))


# Статистика

total_validated = len(good) + len(warning) + len(bad)

print("\nСтатистика валидации")
print(f"  Всего записей:  {total_validated}")

if total_validated > 0:
    print(f"  Good:           {len(good)}  ({len(good)/total_validated:.1%})")
    print(f"  Warning:        {len(warning)}  ({len(warning)/total_validated:.1%})")
    print(f"  Bad:            {len(bad)}  ({len(bad)/total_validated:.1%})")
else:
    print(f"  Good:           {len(good)}")
    print(f"  Warning:        {len(warning)}")
    print(f"  Bad:            {len(bad)}")
    print("  Нет данных для статистики")

if bad:
    print("\nПримеры ошибок (первые 3):")
    for i, (raw, errors) in enumerate(bad[:3], 1):
        print(f"  {i}. Товар: {raw.get('name', 'БЕЗ НАЗВАНИЯ')[:50]}")
        for err in errors:
            field = err["loc"][0] if err["loc"] else "?"
            print(f"     [{field}] {err['msg']}")

if warning:
    print("\nПримеры предупреждений (первые 3):")
    for i, prod in enumerate(warning[:3], 1):
        print(f"  {i}. {prod.name[:50]}")
        for w in prod.warnings:
            print(f"     {w}")

if bad:
    error_types = Counter()
    for _, errors in bad:
        for err in errors:
            error_types[err["type"]] += 1
    print("\nТоп типов ошибок:")
    for err_type, count in error_types.most_common(5):
        print(f"  {err_type}: {count}")


# PostgreSQL

print("\nЗагрузка в PostgreSQL...")

valid_for_db = [p.model_dump(mode="json") for p in good + warning]

for p in valid_for_db:
    brand = p.pop("brand", {})
    p["brand_name"] = brand.get("name", "")
    p["brand_country"] = brand.get("country", "")

if valid_for_db:
    inserted = db.insert_products(valid_for_db)
    print(f"Вставлено в БД: {inserted} строк")
else:
    inserted = 0
    print("Нечего вставлять в БД")


# Excel

print("\nСохранение в Excel...")

wb_good = Workbook()
ws_good = wb_good.active
ws_good.title = "Good"
ws_good.append(["Категория", "Название", "Цена", "Старая цена",
                "Скидка", "Ссылка", "Картинка", "Бренд"])

for p in good:
    ws_good.append([
        p.category.value, p.name, p.price, p.old_price or "",
        p.discount, p.link, p.image_url, p.brand.name,
    ])
wb_good.save("good.xlsx")
print(f"  good.xlsx — {len(good)} строк")

wb_bad = Workbook()
ws_bad = wb_bad.active
ws_bad.title = "Bad"
ws_bad.append(["Raw данные", "Поле", "Ошибка", "Тип"])

for raw, errors in bad:
    for err in errors:
        ws_bad.append([
            str(raw)[:500],
            err["loc"][0] if err["loc"] else "?",
            err["msg"],
            err["type"],
        ])
wb_bad.save("bad.xlsx")
print(f"  bad.xlsx — {len(bad)} строк")

if warning:
    wb_warn = Workbook()
    ws_warn = wb_warn.active
    ws_warn.title = "Warnings"
    ws_warn.append(["Название", "Цена", "Warnings"])
    for p in warning:
        ws_warn.append([p.name, p.price, "; ".join(p.warnings)])
    wb_warn.save("warning.xlsx")
    print(f"  warning.xlsx — {len(warning)} строк")


# JSON

with open("bad_records.json", "w", encoding="utf-8") as f:
    json.dump(
        [
            {
                "raw": {k: str(v) for k, v in raw.items()},
                "errors": [
                    {"field": e["loc"][0] if e["loc"] else "?", "msg": e["msg"], "type": e["type"]}
                    for e in errors
                ],
                "timestamp": datetime.now().isoformat(),
            }
            for raw, errors in bad
        ],
        f,
        ensure_ascii=False,
        indent=2,
    )
print("  bad_records.json — детали ошибок")

stats = {
    "parsed_at": datetime.now().isoformat(),
    "total": total_validated,
    "good": len(good),
    "warning": len(warning),
    "bad": len(bad),
    "success_rate": round((len(good) + len(warning)) / total_validated, 4) if total_validated > 0 else 0,
    "inserted_to_db": inserted,
}

with open("stats.json", "w", encoding="utf-8") as f:
    json.dump(stats, f, ensure_ascii=False, indent=2)

print("  stats.json — сводка")
print("\nГотово.")